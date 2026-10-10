import io
import json
import re
import numpy as np
import pandas as pd

# Control characters that break spreadsheet apps (tab, newline, CR are kept)
_CONTROL_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]')


def clean_text(value):
    if isinstance(value, str):
        return _CONTROL_CHARS.sub(' ', value)
    return value


# Text starting with these is read as a formula by spreadsheet apps
_FORMULA_START = ('=', '+', '-', '@')


def csv_safe(value):
    """Prefix an apostrophe so a spreadsheet shows the text instead of running it as a formula."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def embedding_to_str(x, decimals=8):
    if isinstance(x, np.ndarray) and x.size > 0:
        return json.dumps([round(float(v), decimals) for v in x])
    return ''


def to_csv_bytes(df):
    """CSV with UTF-8 BOM so Excel opens accents/emoji correctly."""
    clean = df.copy()
    for col in clean.columns:
        if not pd.api.types.is_numeric_dtype(clean[col]):
            clean[col] = clean[col].map(clean_text)
    return clean.to_csv(index=False).encode('utf-8-sig')


def embeddings_csv(df, has_categories):
    """URL, Content, (Category), Embedding – can be re-uploaded to skip API costs."""
    return to_csv_bytes(embeddings_frame(df, has_categories))


def results_csv(results_df):
    """Results only hold URLs, labels and numbers, so the formula guard can't alter re-uploaded content."""
    out = results_frame(results_df)
    for col in out.columns:
        if not pd.api.types.is_numeric_dtype(out[col]):
            out[col] = out[col].map(csv_safe)
    return to_csv_bytes(out)


def summary_csv(rows):
    return to_csv_bytes(pd.DataFrame(rows, columns=['Metric', 'Value']))


# ----------------------------------------------------------------- Excel (one file, several sheets)
EXCEL_CELL_LIMIT = 32767


def count_over_cell_limit(series):
    """Number of cells that will be cut to Excel's limit."""
    return int(series.astype(str).str.len().gt(EXCEL_CELL_LIMIT).sum())


def looks_cut_by_excel(value):
    """True for a cell we trimmed to the limit in an earlier export."""
    return isinstance(value, str) and len(value) >= EXCEL_CELL_LIMIT - 100 and value.endswith('…')


def _excel_safe(value):
    if isinstance(value, str):
        value = clean_text(value)
        if len(value) > EXCEL_CELL_LIMIT:
            value = value[:EXCEL_CELL_LIMIT - 1] + '…'
    return value


def embeddings_frame(df, has_categories):
    out = pd.DataFrame({'URL': df['URL'], 'Content': df['Content']})
    if has_categories:
        out['Category'] = df['Category']
    if 'Unique Content %' in df.columns:
        out['Unique Content %'] = df['Unique Content %']
    if 'Page Note' in df.columns:
        out['Page Note'] = df['Page Note']
    out['Embedding'] = df['Embedding'].apply(embedding_to_str)
    if 'Embedding Settings' in df.columns:
        out['Embedding Settings'] = df['Embedding Settings']
    return out


_PERCENT_COLUMNS = ('Semantic Similarity', 'TF-IDF Similarity', 'Full-page Overlap')


def results_frame(results_df):
    out = results_df.copy()
    for col in _PERCENT_COLUMNS:
        if col in out.columns:
            out[col] = (out[col].astype(float) * 100).round(2)
    return out.rename(columns={c: f"{c} (%)" for c in _PERCENT_COLUMNS})


def excel_workbook(sheets):
    """sheets: list of (sheet_name, DataFrame). Returns .xlsx bytes.

    Control characters are removed and over-long cells are trimmed to Excel's
    32,767-character limit, so building the file cannot fail on page content.
    Text starting with '=' is stored as text, not as a formula, without adding
    any characters – the file is meant to be re-uploaded unchanged.
    """
    from openpyxl.styles import Font
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        for name, frame in sheets:
            safe = frame.copy()
            for col in safe.columns:
                if not pd.api.types.is_numeric_dtype(safe[col]):
                    safe[col] = safe[col].map(_excel_safe)
            safe.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = 'A2'
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == 'f':
                        cell.data_type = 's'
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(safe.columns, start=1):
                sample = [str(col)] + [str(v) for v in safe[col].head(50).tolist()]
                width = min(max(len(s) for s in sample) + 2, 60)
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    return buf.getvalue()
